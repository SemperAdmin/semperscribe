"""Regenerates the Batch Generate .xlsx fixtures.

  python3 tests/fixtures/xlsx/make-fixtures.py   (needs openpyxl and LibreOffice)

page11-roster.xlsx      written by openpyxl: shared strings, a Text-formatted
                        EDIPI with a leading zero, a General (number) EDIPI,
                        an EDIPI that lost its leading zero, a real date cell,
                        a multi-line cell, a blank row, and a second sheet.
counseling-roster.xlsx  written by LibreOffice Calc from a CSV, so its
                        styles, shared strings and date cells are what an
                        office suite produces, not what this script chose.
"""
import csv, datetime, os, subprocess, tempfile
from openpyxl import Workbook

here = os.path.dirname(os.path.abspath(__file__))

wb = Workbook()
ws = wb.active
ws.title = 'Roster'
ws.append(['NAME', 'EDIPI', 'DATE', 'ENTRY'])
ws.append(['DOE, JOHN M.', 1234567890, datetime.date(2026, 3, 1), 'First paragraph.\n\nSecond paragraph.'])
ws.append([])
ws.append(['SMITH, JANE A.', '0987654321', datetime.date(2026, 3, 2), 'Short.'])
ws.append(['ROE, RICHARD', 987654321, datetime.date(2026, 12, 31), ''])
ws['B4'].number_format = '@'
for cell in ('C2', 'C4', 'C5'):
    ws[cell].number_format = 'd mmm yy'
notes = wb.create_sheet('Notes')
notes.append(['NOT', 'THE', 'ROSTER'])
wb.save(os.path.join(here, 'page11-roster.xlsx'))

with tempfile.TemporaryDirectory() as tmp:
    src = os.path.join(tmp, 'counseling-roster.csv')
    with open(src, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['LAST_NAME', 'FIRST_NAME', 'MI', 'GRADE', 'EDIPI', 'DOR'])
        w.writerow(['DOE', 'JOHN', 'M', 'LCpl', '1234567890', '06/01/2026'])
        w.writerow(['SMITH', 'JANE', 'A', 'Sgt', '1112223334', '03/01/2025'])
    subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', here, src],
                   check=True, capture_output=True)
